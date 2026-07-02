# thesis/health_rl/lint_deck_layout.py
"""Layout lint for Poly_defense_presentation_v3.pptx (spec 2026-07-02 §8).

Opens the built deck and reports, per slide:
  1. OFF-CANVAS   — a shape spills past the slide edges (tol 15,000 EMU).
                    This is the A5-style bug class. HARD FAIL.
  2. TEXT-OVERLAP — two text-bearing shapes whose bounding boxes intersect
                    by > 15% of the smaller box. HARD FAIL. Only shapes that
                    actually carry non-empty text are compared, so background
                    rects / pictures / empty frames are excluded automatically.
  3. OVERFLOW     — rough estimate that a text frame's content is taller than
                    its box (> 115%). ADVISORY only (heuristic is noisy for
                    vertically-anchored short text in tall boxes).

Exit non-zero if any OFF-CANVAS or TEXT-OVERLAP finding survives.
Usage:  python thesis/health_rl/lint_deck_layout.py
"""
import sys, os
from pptx import Presentation
from pptx.util import Emu

DECK = os.path.join(os.path.dirname(__file__), "Poly_defense_presentation_v3.pptx")
SW = 12191695
SH = 6858000
TOL = 15000            # off-canvas tolerance (EMU)
OVERLAP_FRAC = 0.15    # min intersection / smaller-box area to flag
OVERFLOW_FRAC = 1.15   # advisory overflow threshold


def _geom(shp):
    try:
        l, t, w, h = shp.left, shp.top, shp.width, shp.height
    except Exception:
        return None
    if None in (l, t, w, h):
        return None
    return int(l), int(t), int(w), int(h)


def _text(shp):
    if not getattr(shp, "has_text_frame", False):
        return ""
    return shp.text_frame.text.strip()


def _label(shp):
    txt = _text(shp)
    snip = (txt[:24] + "…") if len(txt) > 24 else txt
    snip = snip.replace("\n", " ")
    return f"{shp.shape_type}:'{snip}'" if snip else f"{shp.shape_type}:{shp.name}"


def _intersect_area(a, b):
    al, at, aw, ah = a
    bl, bt, bw, bh = b
    ix = max(0, min(al + aw, bl + bw) - max(al, bl))
    iy = max(0, min(at + ah, bt + bh) - max(at, bt))
    return ix * iy


def _est_overflow(shp, geom):
    """Very rough: estimate rendered text height vs frame height."""
    tf = shp.text_frame
    _, _, w, h = geom
    w_pt = w / 12700.0
    total_lines = 0
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs) or p.text
        size = None
        for r in p.runs:
            if r.font.size is not None:
                size = r.font.size.pt
                break
        size = size or 12
        cpl = max(1, int(w_pt / (0.50 * size)))   # ~0.5*pt char width
        total_lines += max(1, -(-len(txt) // cpl))
    # dominant font size ~ first run size or 12
    size0 = 12
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                size0 = r.font.size.pt
                break
        else:
            continue
        break
    rendered = total_lines * size0 * 1.2 * 12700
    return rendered, h


def lint():
    prs = Presentation(DECK)
    off_canvas, overlaps, overflow = [], [], []
    for si, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        # 1) off-canvas
        for shp in shapes:
            g = _geom(shp)
            if not g:
                continue
            l, t, w, h = g
            if l < -TOL or t < -TOL or (l + w) > SW + TOL or (t + h) > SH + TOL:
                over = []
                if l < -TOL: over.append(f"left {l}")
                if t < -TOL: over.append(f"top {t}")
                if (l + w) > SW + TOL: over.append(f"right {l+w} > {SW}")
                if (t + h) > SH + TOL: over.append(f"bottom {t+h} > {SH}")
                off_canvas.append((si, _label(shp), ", ".join(over)))
        # 2) text-overlap (only non-empty text-bearing shapes)
        txt_shapes = [(shp, _geom(shp)) for shp in shapes
                      if _text(shp) and _geom(shp)]
        for i in range(len(txt_shapes)):
            for j in range(i + 1, len(txt_shapes)):
                (s1, g1), (s2, g2) = txt_shapes[i], txt_shapes[j]
                inter = _intersect_area(g1, g2)
                if inter <= 0:
                    continue
                smaller = min(g1[2] * g1[3], g2[2] * g2[3])
                if smaller and inter / smaller > OVERLAP_FRAC:
                    overlaps.append((si, _label(s1), _label(s2),
                                     f"{inter/smaller:.0%} of smaller"))
        # 3) overflow (advisory)
        for shp in shapes:
            if not _text(shp):
                continue
            g = _geom(shp)
            if not g:
                continue
            rendered, box_h = _est_overflow(shp, g)
            if box_h and rendered > OVERFLOW_FRAC * box_h:
                overflow.append((si, _label(shp),
                                 f"~{rendered/box_h:.0%} of box height"))

    print(f"=== Layout lint: {len(prs.slides)} slides ===")

    def _report(title, rows, fmt):
        print(f"\n## {title}: {len(rows)}")
        for r in rows:
            print("  slide %2d  %s" % (r[0], fmt(r)))

    _report("OFF-CANVAS (hard)", off_canvas, lambda r: f"{r[1]}  [{r[2]}]")
    _report("TEXT-OVERLAP (hard)", overlaps,
            lambda r: f"{r[1]}  ×  {r[2]}   ({r[3]})")
    _report("OVERFLOW (advisory)", overflow, lambda r: f"{r[1]}  [{r[2]}]")

    hard = len(off_canvas) + len(overlaps)
    print(f"\n=== {hard} hard finding(s), {len(overflow)} advisory ===")
    return 1 if hard else 0


if __name__ == "__main__":
    sys.exit(lint())

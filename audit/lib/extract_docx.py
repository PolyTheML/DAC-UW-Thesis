import re
from pathlib import Path
from docx import Document
from lxml import etree

_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_M = f"{{{_M_NS}}}"

_DISPLAY_RE = re.compile(r'\$\$(.*?)\$\$', re.DOTALL)
_INLINE_RE = re.compile(r'(?<!\$)\$(?!\$)((?:[^$]|\n)+?)(?<!\$)\$(?!\$)')


# ── OMML → LaTeX converter (limited to constructs in this thesis) ─────────────

def _omml_text(node) -> str:
    """Recursively extract text from an OMML node."""
    parts = []
    for child in node:
        tag = child.tag.replace(_M, "m:")
        if tag == "m:r":
            for sub in child:
                if "t" in sub.tag:
                    parts.append(sub.text or "")
        elif tag in ("m:f", "m:rad", "m:sup", "m:sub", "m:nary", "m:func",
                     "m:e", "m:num", "m:den", "m:deg", "m:fName"):
            parts.append(_omml_node_to_latex(child))
        else:
            parts.append(_omml_text(child))
    return "".join(parts)


def _omml_node_to_latex(node) -> str:
    tag = node.tag.replace(_M, "m:")

    if tag == "m:r":
        return "".join(sub.text or "" for sub in node if "t" in sub.tag)

    if tag == "m:f":
        num = node.find(f"{_M}num")
        den = node.find(f"{_M}den")
        n = _omml_text(num) if num is not None else ""
        d = _omml_text(den) if den is not None else ""
        return f"\\frac{{{n}}}{{{d}}}"

    if tag == "m:sup":
        base = node.find(f"{_M}e")
        sup = node.find(f"{_M}sup")
        b = _omml_text(base) if base is not None else ""
        s = _omml_text(sup) if sup is not None else ""
        return f"{b}^{{{s}}}"

    if tag == "m:sub":
        base = node.find(f"{_M}e")
        sub = node.find(f"{_M}sub")
        b = _omml_text(base) if base is not None else ""
        s = _omml_text(sub) if sub is not None else ""
        return f"{b}_{{{s}}}"

    if tag == "m:rad":
        e = node.find(f"{_M}e")
        e_tex = _omml_text(e) if e is not None else ""
        return f"\\sqrt{{{e_tex}}}"

    if tag == "m:nary":
        pr = node.find(f"{_M}naryPr")
        chr_el = pr.find(f"{_M}chr") if pr is not None else None
        op_char = chr_el.get(f"{{{_M_NS}}}val", "∑") if chr_el is not None else "∑"
        op = {"∑": "\\sum", "∏": "\\prod", "∫": "\\int"}.get(op_char, "\\sum")
        sub_node = node.find(f"{_M}sub")
        sup_node = node.find(f"{_M}sup")
        e = node.find(f"{_M}e")
        result = op
        if sub_node is not None:
            result += f"_{{{_omml_text(sub_node)}}}"
        if sup_node is not None:
            result += f"^{{{_omml_text(sup_node)}}}"
        if e is not None:
            result += " " + _omml_text(e)
        return result

    # Default: recurse
    return _omml_text(node)


def _extract_omml_equations(doc: Document) -> list[dict]:
    """Extract equations stored as OMML (native Word math objects)."""
    equations = []
    for para_idx, para in enumerate(doc.paragraphs):
        try:
            xml = etree.fromstring(para._element.xml)
        except (etree.XMLSyntaxError, AttributeError):
            continue
        math_nodes = xml.findall(f".//{_M}oMath")
        if not math_nodes:
            math_nodes = xml.findall(f".//{_M}oMathPara//{_M}oMath")
        for node in math_nodes:
            latex = _omml_node_to_latex(node).strip()
            if latex:
                equations.append({
                    "latex_docx": latex,
                    "para_idx": para_idx,
                    "method": "omml",
                })
    return equations


def _extract_text_equations(doc: Document) -> list[dict]:
    """Fallback: extract $$...$$ and $...$ from paragraph text."""
    equations = []
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text)

    for m in _DISPLAY_RE.finditer(full_text):
        equations.append({
            "latex_docx": m.group(1).strip(),
            "para_idx": -1,
            "method": "text_display",
        })

    display_spans = [(m.start(), m.end()) for m in _DISPLAY_RE.finditer(full_text)]

    def in_display(pos: int) -> bool:
        return any(s <= pos <= e for s, e in display_spans)

    for m in _INLINE_RE.finditer(full_text):
        latex = m.group(1).strip()
        if not in_display(m.start()) and len(latex) >= 2:
            # Skip currency amounts (same filter as extract_md)
            if latex[0].isdigit() and not any(c in latex for c in r'\_{^'):
                continue
            equations.append({
                "latex_docx": latex,
                "para_idx": -1,
                "method": "text_inline",
            })
    return equations


def extract_equations_from_docx(docx_path: Path) -> list[dict]:
    """
    Extract equations from a DOCX file.
    Tries OMML first; falls back to text regex if no OMML nodes found.
    Returns list of dicts with keys: latex_docx, para_idx, method.
    """
    doc = Document(str(docx_path))
    equations = _extract_omml_equations(doc)
    if not equations:
        equations = _extract_text_equations(doc)
    return equations
